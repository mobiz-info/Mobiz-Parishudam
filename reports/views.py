from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.models import Branch
from operations.models import DailySale, Expense


# =========================================================
# DAILY P&L REPORT
# =========================================================

@login_required
def report_daily_view(request):
    user_profile = getattr(request.user, 'profile', None)
    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    date_str = request.GET.get('date')

    if date_str:
        try:
            report_date = datetime.strptime(
                date_str,
                '%Y-%m-%d'
            ).date()
        except ValueError:
            report_date = timezone.now().date()
    else:
        report_date = timezone.now().date()

    # -----------------------------------------------------
    # Branch Selection
    # -----------------------------------------------------

    if not is_admin and user_profile and user_profile.branch:

        selected_branch = user_profile.branch

    else:

        branch_id = request.GET.get('branch_id')

        selected_branch = (
            Branch.objects.filter(
                id=branch_id
            ).first()
            if branch_id
            else Branch.objects.filter(
                status='active'
            ).first()
        )

    # -----------------------------------------------------
    # Querysets
    # -----------------------------------------------------

    sales = DailySale.objects.filter(
        sale_date=report_date
    )

    expenses = Expense.objects.filter(
        expense_date=report_date
    )

    if selected_branch:

        sales = sales.filter(
            branch=selected_branch
        )

        expenses = expenses.filter(
            branch=selected_branch
        )

    # -----------------------------------------------------
    # Details
    # -----------------------------------------------------

    sales_detail = sales.select_related(
        'product',
        'product_packing'
    )

    expenses_detail = expenses.select_related(
        'expense_head',
        'staff'
    )

    # -----------------------------------------------------
    # Totals
    # -----------------------------------------------------

    gross_profit = sales.aggregate(
        val=Sum('profit')
    )['val'] or 0.00

    total_expense = expenses.aggregate(
        val=Sum('amount')
    )['val'] or 0.00

    net_profit = (
        float(gross_profit)
        - float(total_expense)
    )

    total_volume_litres = sales.aggregate(
        val=Sum('base_quantity')
    )['val'] or 0.0000

    # -----------------------------------------------------
    # Branches
    # -----------------------------------------------------

    branches = Branch.objects.filter(
        status='active'
    )

    # -----------------------------------------------------
    # Render
    # -----------------------------------------------------

    return render(
        request,
        'reports/report_daily.html',
        {
            'report_date': report_date,
            'selected_branch': selected_branch,
            'branches': branches,
            'sales_detail': sales_detail,
            'expenses_detail': expenses_detail,
            'gross_profit': gross_profit,
            'total_expense': total_expense,
            'net_profit': net_profit,
            'total_volume_litres': total_volume_litres,
            'is_admin': is_admin,
        }
    )


# =========================================================
# WEEKLY P&L REPORT
# =========================================================

@login_required
def report_weekly_view(request):
    user_profile = getattr(request.user, 'profile', None)

    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    today = timezone.now().date()

    # -----------------------------------------------------
    # Branch Selection
    # -----------------------------------------------------

    if not is_admin and user_profile and user_profile.branch:

        selected_branch = user_profile.branch

    else:

        branch_id = request.GET.get('branch_id')

        if branch_id:

            selected_branch = Branch.objects.filter(
                id=branch_id,
                status='active'
            ).first()

        else:

            selected_branch = None

    # -----------------------------------------------------
    # Start Date
    # -----------------------------------------------------

    date_from_str = request.GET.get(
        'date_from'
    )

    if date_from_str:

        try:

            date_from = datetime.strptime(
                date_from_str,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            date_from = today

    else:

        date_from = today

    # -----------------------------------------------------
    # Weekly = 7 Days
    # -----------------------------------------------------

    date_to = date_from + timedelta(
        days=6
    )

    # -----------------------------------------------------
    # Querysets
    # -----------------------------------------------------

    sales = DailySale.objects.filter(
        sale_date__range=[
            date_from,
            date_to
        ]
    )

    expenses = Expense.objects.filter(
        expense_date__range=[
            date_from,
            date_to
        ]
    )

    # -----------------------------------------------------
    # Branch Filter
    # -----------------------------------------------------

    if selected_branch:

        sales = sales.filter(
            branch=selected_branch
        )

        expenses = expenses.filter(
            branch=selected_branch
        )

    # -----------------------------------------------------
    # Overall Totals
    # -----------------------------------------------------

    sale_profit = sales.aggregate(
        total=Sum('profit')
    )['total'] or 0.00

    total_expense = expenses.aggregate(
        total=Sum('amount')
    )['total'] or 0.00

    total_sold_volume = sales.aggregate(
        total=Sum('base_quantity')
    )['total'] or 0.00

    net_profit = (
        float(sale_profit)
        - float(total_expense)
    )

    # -----------------------------------------------------
    # Daily Rows
    # -----------------------------------------------------

    report_rows = []

    current_date = date_from

    while current_date <= date_to:

        day_sales = sales.filter(
            sale_date=current_date
        )

        day_expenses = expenses.filter(
            expense_date=current_date
        )

        day_sale_profit = day_sales.aggregate(
            total=Sum('profit')
        )['total'] or 0.00

        day_expense = day_expenses.aggregate(
            total=Sum('amount')
        )['total'] or 0.00

        day_net_profit = (
            float(day_sale_profit)
            - float(day_expense)
        )

        report_rows.append({
            'date': current_date,
            'period': current_date.strftime(
                '%d-%b-%Y'
            ),
            'day': current_date.strftime(
                '%A'
            ),
            'sale_profit': day_sale_profit,
            'expense': day_expense,
            'net_profit': day_net_profit,
        })

        current_date += timedelta(
            days=1
        )

    # -----------------------------------------------------
    # Branches
    # -----------------------------------------------------

    branches = Branch.objects.filter(
        status='active'
    )

    # -----------------------------------------------------
    # Render
    # -----------------------------------------------------

    return render(
        request,
        'reports/report_weekly.html',
        {
            'date_from': date_from,
            'date_to': date_to,
            'selected_branch': selected_branch,
            'branches': branches,
            'sale_profit': sale_profit,
            'total_expense': total_expense,
            'total_sold_volume': total_sold_volume,
            'net_profit': net_profit,
            'report_rows': report_rows,
            'is_admin': is_admin,
        }
    )


# =========================================================
# MONTHLY / YEARLY P&L REPORT
# =========================================================

@login_required
def report_monthly_yearly_view(request):
    user_profile = getattr(request.user, 'profile', None)

    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    today = timezone.now().date()

    # -----------------------------------------------------
    # Report Type
    # -----------------------------------------------------

    report_type = request.GET.get(
        'report_type',
        'monthly'
    )

    if report_type not in [
        'monthly',
        'yearly'
    ]:
        report_type = 'monthly'

    # -----------------------------------------------------
    # Branch Selection
    # -----------------------------------------------------

    if not is_admin and user_profile and user_profile.branch:

        selected_branch = user_profile.branch

    else:

        branch_id = request.GET.get(
            'branch_id'
        )

        if branch_id:

            selected_branch = Branch.objects.filter(
                id=branch_id,
                status='active'
            ).first()

        else:

            selected_branch = None

    # -----------------------------------------------------
    # Default Values
    # -----------------------------------------------------

    selected_month = today.strftime(
        '%Y-%m'
    )

    selected_year = str(
        today.year
    )

    # -----------------------------------------------------
    # Monthly
    # -----------------------------------------------------

    if report_type == 'monthly':

        month_str = request.GET.get(
            'month'
        )

        if month_str:

            try:

                selected_month_date = datetime.strptime(
                    month_str,
                    '%Y-%m'
                ).date()

                selected_month = month_str

            except ValueError:

                selected_month_date = today

                selected_month = today.strftime(
                    '%Y-%m'
                )

        else:

            selected_month_date = today

        date_from = selected_month_date.replace(
            day=1
        )

        if date_from.month == 12:

            next_month = date_from.replace(
                year=date_from.year + 1,
                month=1,
                day=1
            )

        else:

            next_month = date_from.replace(
                month=date_from.month + 1,
                day=1
            )

        date_to = next_month - timedelta(
            days=1
        )

    # -----------------------------------------------------
    # Yearly
    # -----------------------------------------------------

    else:

        year_str = request.GET.get(
            'year',
            str(today.year)
        )

        try:

            year_value = int(
                year_str
            )

        except (ValueError, TypeError):

            year_value = today.year

        selected_year = str(
            year_value
        )

        date_from = datetime(
            year_value,
            1,
            1
        ).date()

        date_to = datetime(
            year_value,
            12,
            31
        ).date()

    # -----------------------------------------------------
    # Querysets
    # -----------------------------------------------------

    sales = DailySale.objects.filter(
        sale_date__range=[
            date_from,
            date_to
        ]
    )

    expenses = Expense.objects.filter(
        expense_date__range=[
            date_from,
            date_to
        ]
    )

    # -----------------------------------------------------
    # Branch Filter
    # -----------------------------------------------------

    if selected_branch:

        sales = sales.filter(
            branch=selected_branch
        )

        expenses = expenses.filter(
            branch=selected_branch
        )

    # -----------------------------------------------------
    # Totals
    # -----------------------------------------------------

    sale_profit = sales.aggregate(
        total=Sum('profit')
    )['total'] or 0.00

    total_expense = expenses.aggregate(
        total=Sum('amount')
    )['total'] or 0.00

    total_sold_volume = sales.aggregate(
        total=Sum('base_quantity')
    )['total'] or 0.00

    net_profit = (
        float(sale_profit)
        - float(total_expense)
    )

    # -----------------------------------------------------
    # Report Rows
    # -----------------------------------------------------

    report_rows = []

    # -----------------------------------------------------
    # MONTHLY = DAY WISE
    # -----------------------------------------------------

    if report_type == 'monthly':

        current_date = date_from

        while current_date <= date_to:

            day_sales = sales.filter(
                sale_date=current_date
            )

            day_expenses = expenses.filter(
                expense_date=current_date
            )

            day_sale_profit = day_sales.aggregate(
                total=Sum('profit')
            )['total'] or 0.00

            day_expense = day_expenses.aggregate(
                total=Sum('amount')
            )['total'] or 0.00

            day_net_profit = (
                float(day_sale_profit)
                - float(day_expense)
            )

            report_rows.append({
                'date': current_date,
                'period': current_date.strftime(
                    '%d-%b-%Y'
                ),
                'day': current_date.strftime(
                    '%A'
                ),
                'sale_profit': day_sale_profit,
                'expense': day_expense,
                'net_profit': day_net_profit,
            })

            current_date += timedelta(
                days=1
            )

    # -----------------------------------------------------
    # YEARLY = MONTH WISE
    # -----------------------------------------------------

    else:

        current_date = date_from

        while current_date <= date_to:

            if current_date.month == 12:

                next_month = current_date.replace(
                    year=current_date.year + 1,
                    month=1,
                    day=1
                )

            else:

                next_month = current_date.replace(
                    month=current_date.month + 1,
                    day=1
                )

            month_end = next_month - timedelta(
                days=1
            )

            if month_end > date_to:

                month_end = date_to

            month_sales = sales.filter(
                sale_date__range=[
                    current_date,
                    month_end
                ]
            )

            month_expenses = expenses.filter(
                expense_date__range=[
                    current_date,
                    month_end
                ]
            )

            month_sale_profit = month_sales.aggregate(
                total=Sum('profit')
            )['total'] or 0.00

            month_expense = month_expenses.aggregate(
                total=Sum('amount')
            )['total'] or 0.00

            month_net_profit = (
                float(month_sale_profit)
                - float(month_expense)
            )

            report_rows.append({
                'date': current_date,
                'period': current_date.strftime(
                    '%B'
                ),
                'day': current_date.strftime(
                    '%Y'
                ),
                'sale_profit': month_sale_profit,
                'expense': month_expense,
                'net_profit': month_net_profit,
            })

            current_date = next_month

    # -----------------------------------------------------
    # Branches
    # -----------------------------------------------------

    branches = Branch.objects.filter(
        status='active'
    )

    # -----------------------------------------------------
    # Years
    # -----------------------------------------------------

    years = range(
        today.year - 5,
        today.year + 6
    )

    # -----------------------------------------------------
    # Render
    # -----------------------------------------------------

    return render(
        request,
        'reports/report_monthly_yearly.html',
        {
            'report_type': report_type,
            'date_from': date_from,
            'date_to': date_to,
            'selected_month': selected_month,
            'selected_year': selected_year,
            'selected_branch': selected_branch,
            'branches': branches,
            'sale_profit': sale_profit,
            'total_expense': total_expense,
            'total_sold_volume': total_sold_volume,
            'net_profit': net_profit,
            'report_rows': report_rows,
            'is_admin': is_admin,
            'years': years,
        }
    )


# =========================================================
# DAILY EXCEL EXPORT
# =========================================================

@login_required
def report_daily_export_excel_view(request):
    user_profile = getattr(request.user, 'profile', None)

    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    date_str = request.GET.get(
        'date'
    )

    if date_str:

        try:

            report_date = datetime.strptime(
                date_str,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            report_date = timezone.now().date()

    else:

        report_date = timezone.now().date()

    # -----------------------------------------------------
    # Branch
    # -----------------------------------------------------

    if not is_admin and user_profile and user_profile.branch:

        selected_branch = user_profile.branch

    else:

        branch_id = request.GET.get(
            'branch_id'
        )

        selected_branch = (
            Branch.objects.filter(
                id=branch_id,
                status='active'
            ).first()
            if branch_id
            else Branch.objects.filter(
                status='active'
            ).first()
        )

    # -----------------------------------------------------
    # Querysets
    # -----------------------------------------------------

    sales = DailySale.objects.filter(
        sale_date=report_date
    )

    expenses = Expense.objects.filter(
        expense_date=report_date
    )

    if selected_branch:

        sales = sales.filter(
            branch=selected_branch
        )

        expenses = expenses.filter(
            branch=selected_branch
        )

    # -----------------------------------------------------
    # Workbook
    # -----------------------------------------------------

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Daily P&L'

    sheet.append([
        'Date',
        'Sale Profit',
        'Expense',
        'Net Profit'
    ])

    gross_profit = sales.aggregate(
        total=Sum('profit')
    )['total'] or 0.00

    total_expense = expenses.aggregate(
        total=Sum('amount')
    )['total'] or 0.00

    net_profit = (
        float(gross_profit)
        - float(total_expense)
    )

    sheet.append([
        report_date,
        float(gross_profit),
        float(total_expense),
        float(net_profit)
    ])

    # -----------------------------------------------------
    # Total Row
    # -----------------------------------------------------

    sheet.append([
        'TOTAL',
        float(gross_profit),
        float(total_expense),
        float(net_profit)
    ])

    total_row = sheet.max_row

    # -----------------------------------------------------
    # Header Style
    # -----------------------------------------------------

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

    # -----------------------------------------------------
    # Total Style
    # -----------------------------------------------------

    for cell in sheet[total_row]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        cell.fill = PatternFill(
            fill_type='solid',
            fgColor='E0E7FF'
        )

    # -----------------------------------------------------
    # Date Format
    # -----------------------------------------------------

    for cell in sheet['A'][1:]:

        if hasattr(
            cell.value,
            'strftime'
        ):

            cell.number_format = (
                'DD-MM-YYYY'
            )

    # -----------------------------------------------------
    # Alignment
    # -----------------------------------------------------

    for row in sheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                horizontal='center',
                vertical='center'
            )

    # -----------------------------------------------------
    # Column Widths
    # -----------------------------------------------------

    sheet.column_dimensions['A'].width = 18
    sheet.column_dimensions['B'].width = 18
    sheet.column_dimensions['C'].width = 18
    sheet.column_dimensions['D'].width = 18

    # -----------------------------------------------------
    # Response
    # -----------------------------------------------------

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="daily_profit_loss.xlsx"'
    )

    workbook.save(
        response
    )

    return response


# =========================================================
# WEEKLY EXCEL EXPORT
# =========================================================

@login_required
def report_weekly_export_excel_view(request):
    user_profile = getattr(request.user, 'profile', None)

    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    today = timezone.now().date()

    # -----------------------------------------------------
    # Branch
    # -----------------------------------------------------

    if not is_admin and user_profile and user_profile.branch:

        selected_branch = user_profile.branch

    else:

        branch_id = request.GET.get(
            'branch_id'
        )

        selected_branch = (
            Branch.objects.filter(
                id=branch_id,
                status='active'
            ).first()
            if branch_id
            else None
        )

    # -----------------------------------------------------
    # Date
    # -----------------------------------------------------

    date_from_str = request.GET.get(
        'date_from'
    )

    if date_from_str:

        try:

            date_from = datetime.strptime(
                date_from_str,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            date_from = today

    else:

        date_from = today

    date_to = date_from + timedelta(
        days=6
    )

    # -----------------------------------------------------
    # Querysets
    # -----------------------------------------------------

    sales = DailySale.objects.filter(
        sale_date__range=[
            date_from,
            date_to
        ]
    )

    expenses = Expense.objects.filter(
        expense_date__range=[
            date_from,
            date_to
        ]
    )

    if selected_branch:

        sales = sales.filter(
            branch=selected_branch
        )

        expenses = expenses.filter(
            branch=selected_branch
        )

    # -----------------------------------------------------
    # Totals
    # -----------------------------------------------------

    sale_profit = sales.aggregate(
        total=Sum('profit')
    )['total'] or 0.00

    total_expense = expenses.aggregate(
        total=Sum('amount')
    )['total'] or 0.00

    net_profit = (
        float(sale_profit)
        - float(total_expense)
    )

    # -----------------------------------------------------
    # Workbook
    # -----------------------------------------------------

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Weekly P&L'

    sheet.append([
        'Date',
        'Day',
        'Sale Profit',
        'Expense',
        'Net Profit'
    ])

    current_date = date_from

    while current_date <= date_to:

        day_sales = sales.filter(
            sale_date=current_date
        )

        day_expenses = expenses.filter(
            expense_date=current_date
        )

        day_sale_profit = day_sales.aggregate(
            total=Sum('profit')
        )['total'] or 0.00

        day_expense = day_expenses.aggregate(
            total=Sum('amount')
        )['total'] or 0.00

        day_net_profit = (
            float(day_sale_profit)
            - float(day_expense)
        )

        sheet.append([
            current_date,
            current_date.strftime('%A'),
            float(day_sale_profit),
            float(day_expense),
            float(day_net_profit)
        ])

        current_date += timedelta(
            days=1
        )

    # -----------------------------------------------------
    # Total Row
    # -----------------------------------------------------

    sheet.append([
        'TOTAL',
        '',
        float(sale_profit),
        float(total_expense),
        float(net_profit)
    ])

    total_row = sheet.max_row

    # -----------------------------------------------------
    # Header
    # -----------------------------------------------------

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

    # -----------------------------------------------------
    # Total
    # -----------------------------------------------------

    for cell in sheet[total_row]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        cell.fill = PatternFill(
            fill_type='solid',
            fgColor='E0E7FF'
        )

    # -----------------------------------------------------
    # Date Format
    # -----------------------------------------------------

    for cell in sheet['A'][1:]:

        if hasattr(
            cell.value,
            'strftime'
        ):

            cell.number_format = (
                'DD-MM-YYYY'
            )

    # -----------------------------------------------------
    # Alignment
    # -----------------------------------------------------

    for row in sheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                horizontal='center',
                vertical='center'
            )

    # -----------------------------------------------------
    # Column Width
    # -----------------------------------------------------

    sheet.column_dimensions['A'].width = 18
    sheet.column_dimensions['B'].width = 18
    sheet.column_dimensions['C'].width = 18
    sheet.column_dimensions['D'].width = 18
    sheet.column_dimensions['E'].width = 18

    # -----------------------------------------------------
    # Response
    # -----------------------------------------------------

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="weekly_profit_loss.xlsx"'
    )

    workbook.save(
        response
    )

    return response


# =========================================================
# MONTHLY / YEARLY EXCEL EXPORT
# =========================================================

@login_required
def report_monthly_yearly_export_excel_view(request):
    user_profile = getattr(request.user, 'profile', None)

    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    today = timezone.now().date()

    # -----------------------------------------------------
    # Report Type
    # -----------------------------------------------------

    report_type = request.GET.get(
        'report_type',
        'monthly'
    )

    if report_type not in [
        'monthly',
        'yearly'
    ]:
        report_type = 'monthly'

    # -----------------------------------------------------
    # Branch
    # -----------------------------------------------------

    if not is_admin and user_profile and user_profile.branch:

        selected_branch = user_profile.branch

    else:

        branch_id = request.GET.get(
            'branch_id'
        )

        selected_branch = (
            Branch.objects.filter(
                id=branch_id,
                status='active'
            ).first()
            if branch_id
            else None
        )

    # -----------------------------------------------------
    # Date Range
    # -----------------------------------------------------

    if report_type == 'monthly':

        month_str = request.GET.get(
            'month'
        )

        try:

            selected_month_date = datetime.strptime(
                month_str,
                '%Y-%m'
            ).date() if month_str else today

        except ValueError:

            selected_month_date = today

        date_from = selected_month_date.replace(
            day=1
        )

        if date_from.month == 12:

            next_month = date_from.replace(
                year=date_from.year + 1,
                month=1,
                day=1
            )

        else:

            next_month = date_from.replace(
                month=date_from.month + 1,
                day=1
            )

        date_to = next_month - timedelta(
            days=1
        )

    else:

        year_str = request.GET.get(
            'year',
            str(today.year)
        )

        try:

            year_value = int(
                year_str
            )

        except (ValueError, TypeError):

            year_value = today.year

        date_from = datetime(
            year_value,
            1,
            1
        ).date()

        date_to = datetime(
            year_value,
            12,
            31
        ).date()

    # -----------------------------------------------------
    # Querysets
    # -----------------------------------------------------

    sales = DailySale.objects.filter(
        sale_date__range=[
            date_from,
            date_to
        ]
    )

    expenses = Expense.objects.filter(
        expense_date__range=[
            date_from,
            date_to
        ]
    )

    if selected_branch:

        sales = sales.filter(
            branch=selected_branch
        )

        expenses = expenses.filter(
            branch=selected_branch
        )

    # -----------------------------------------------------
    # Totals
    # -----------------------------------------------------

    sale_profit = sales.aggregate(
        total=Sum('profit')
    )['total'] or 0.00

    total_expense = expenses.aggregate(
        total=Sum('amount')
    )['total'] or 0.00

    net_profit = (
        float(sale_profit)
        - float(total_expense)
    )

    # -----------------------------------------------------
    # Workbook
    # -----------------------------------------------------

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'P&L Report'

    sheet.append([
        'Date / Period',
        'Day / Year',
        'Sale Profit',
        'Expense',
        'Net Profit'
    ])

    # -----------------------------------------------------
    # Monthly = Day Wise
    # -----------------------------------------------------

    if report_type == 'monthly':

        current_date = date_from

        while current_date <= date_to:

            day_sales = sales.filter(
                sale_date=current_date
            )

            day_expenses = expenses.filter(
                expense_date=current_date
            )

            day_sale_profit = day_sales.aggregate(
                total=Sum('profit')
            )['total'] or 0.00

            day_expense = day_expenses.aggregate(
                total=Sum('amount')
            )['total'] or 0.00

            day_net_profit = (
                float(day_sale_profit)
                - float(day_expense)
            )

            sheet.append([
                current_date,
                current_date.strftime('%A'),
                float(day_sale_profit),
                float(day_expense),
                float(day_net_profit)
            ])

            current_date += timedelta(
                days=1
            )

    # -----------------------------------------------------
    # Yearly = Month Wise
    # -----------------------------------------------------

    else:

        current_date = date_from

        while current_date <= date_to:

            if current_date.month == 12:

                next_month = current_date.replace(
                    year=current_date.year + 1,
                    month=1,
                    day=1
                )

            else:

                next_month = current_date.replace(
                    month=current_date.month + 1,
                    day=1
                )

            month_end = next_month - timedelta(
                days=1
            )

            if month_end > date_to:

                month_end = date_to

            month_sales = sales.filter(
                sale_date__range=[
                    current_date,
                    month_end
                ]
            )

            month_expenses = expenses.filter(
                expense_date__range=[
                    current_date,
                    month_end
                ]
            )

            month_sale_profit = month_sales.aggregate(
                total=Sum('profit')
            )['total'] or 0.00

            month_expense = month_expenses.aggregate(
                total=Sum('amount')
            )['total'] or 0.00

            month_net_profit = (
                float(month_sale_profit)
                - float(month_expense)
            )

            sheet.append([
                current_date.strftime('%B'),
                current_date.strftime('%Y'),
                float(month_sale_profit),
                float(month_expense),
                float(month_net_profit)
            ])

            current_date = next_month

    # -----------------------------------------------------
    # Total Row
    # -----------------------------------------------------

    sheet.append([
        'TOTAL',
        '',
        float(sale_profit),
        float(total_expense),
        float(net_profit)
    ])

    total_row = sheet.max_row

    # -----------------------------------------------------
    # Header
    # -----------------------------------------------------

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

    # -----------------------------------------------------
    # Total
    # -----------------------------------------------------

    for cell in sheet[total_row]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        cell.fill = PatternFill(
            fill_type='solid',
            fgColor='E0E7FF'
        )

    # -----------------------------------------------------
    # Date Format
    # -----------------------------------------------------

    for cell in sheet['A'][1:]:

        if hasattr(
            cell.value,
            'strftime'
        ):

            cell.number_format = (
                'DD-MM-YYYY'
            )

    # -----------------------------------------------------
    # Alignment
    # -----------------------------------------------------

    for row in sheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                horizontal='center',
                vertical='center'
            )

    # -----------------------------------------------------
    # Width
    # -----------------------------------------------------

    sheet.column_dimensions['A'].width = 20
    sheet.column_dimensions['B'].width = 18
    sheet.column_dimensions['C'].width = 18
    sheet.column_dimensions['D'].width = 18
    sheet.column_dimensions['E'].width = 18

    # -----------------------------------------------------
    # Response
    # -----------------------------------------------------

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="monthly_yearly_profit_loss.xlsx"'
    )

    workbook.save(
        response
    )

    return response