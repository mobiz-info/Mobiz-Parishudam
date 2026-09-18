from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from datetime import datetime
from django.db.models import Sum
from decimal import Decimal
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from core.models import Branch, StaffProfile, Product, ProductPackingSize, ProductMargin 
from operations.models import DailySale, Expense, ExpenseHead,Vehicle


# --- DAILY SALES MATRIX ENTRY ---
@login_required
def sales_entry_view(request):
    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    # Determine branch
    branch_id = request.GET.get('branch_id')
    if not is_admin and user_profile and user_profile.branch:
        branch = user_profile.branch
    elif branch_id:
        branch = get_object_or_404(Branch, id=branch_id)
    else:
        branch = Branch.objects.filter(status='active').first()

    # Determine date
    date_str = request.GET.get('sale_date')
    if date_str:
        try:
            sale_date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            sale_date = timezone.now().date()
    else:
        sale_date = timezone.now().date()

    branches = Branch.objects.filter(status='active')
    packings = ProductPackingSize.objects.select_related('product', 'packing_unit').all()

    if request.method == 'POST':
        # Process matrix submission
        form_sale_date_str = request.POST.get('sale_date')
        if form_sale_date_str:
            sale_date = datetime.strptime(form_sale_date_str, '%Y-%m-%d').date()

        selected_branch_id = request.POST.get('branch_id')
        if is_admin and selected_branch_id:
            branch = get_object_or_404(Branch, id=selected_branch_id)

        saved_count = 0
        for p in packings:
            count_key = f"quantity_{p.id}"
            margin_key = f"margin_{p.id}"
            
            count_val = request.POST.get(count_key)
            margin_val = request.POST.get(margin_key)

            if count_val is not None and count_val.strip() != "":
                try:
                    count_int = int(count_val)
                    if count_int < 0:
                        count_int = 0
                except ValueError:
                    count_int = 0

                try:
                    if margin_val and margin_val.strip() != "":
                        margin_dec = Decimal(margin_val)
                    else:
                        margin_obj = ProductMargin.objects.filter(
                            product=p.product,
                            packing_size=p,
                            effective_date__lte=sale_date
                        ).order_by('-effective_date').first()
                        margin_dec = margin_obj.margin_amount if margin_obj else Decimal('0.00')
                except Exception:
                    margin_obj = ProductMargin.objects.filter(
                        product=p.product,
                        packing_size=p,
                        effective_date__lte=sale_date
                    ).order_by('-effective_date').first()
                    margin_dec = margin_obj.margin_amount if margin_obj else Decimal('0.00')

                if count_int >= 0:
                    daily_sale, created = DailySale.objects.get_or_create(
                        branch=branch,
                        sale_date=sale_date,
                        product_packing=p,
                        defaults={
                            'staff': request.user,
                            'product': p.product,
                            'packing_count': count_int,
                            'margin': margin_dec,
                            'created_by': request.user,
                            'updated_by': request.user,
                        }
                    )
                    if not created:
                        daily_sale.packing_count = count_int
                        daily_sale.margin = margin_dec
                        daily_sale.updated_by = request.user
                        daily_sale.save()
                    saved_count += 1

        messages.success(request, f"Daily sales for {branch.name} on {sale_date.strftime('%d-%b-%Y')} saved successfully.")
        return redirect(f"/sales/entry/?sale_date={sale_date.strftime('%Y-%m-%d')}&branch_id={branch.id}")

    # Build packing matrix data
    matrix = []
    for p in packings:
        existing_sale = DailySale.objects.filter(branch=branch, sale_date=sale_date, product_packing=p).first()
        if existing_sale:
            effective_margin = existing_sale.margin
        else:
            margin_obj = ProductMargin.objects.filter(
                product=p.product,
                packing_size=p,
                effective_date__lte=sale_date
            ).order_by('-effective_date').first()
            effective_margin = margin_obj.margin_amount if margin_obj else Decimal('0.00')

        qty_count = existing_sale.packing_count if existing_sale else 0
        profit_calc = qty_count * float(effective_margin)
        base_litres = qty_count * float(p.base_qty_unit)

        matrix.append({
            'packing': p,
            'margin': effective_margin,
            'count': qty_count,
            'profit': profit_calc,
            'base_litres': base_litres
        })

    context = {
        'branch': branch,
        'branches': branches,
        'sale_date': sale_date,
        'matrix': matrix,
        'is_admin': is_admin,
    }
    return render(request, 'operations/sales_entry.html', context)


@login_required
def sales_list_view(request):
    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    sales_qs = DailySale.objects.select_related('branch', 'staff', 'product', 'product_packing').all()

    if not is_admin and user_profile and user_profile.branch:
        sales_qs = sales_qs.filter(branch=user_profile.branch)
    else:
        branch_id = request.GET.get('branch_id')
        if branch_id:
            sales_qs = sales_qs.filter(branch_id=branch_id)

    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    if date_from:
        sales_qs = sales_qs.filter(sale_date__gte=date_from)
    if date_to:
        sales_qs = sales_qs.filter(sale_date__lte=date_to)

    branches = Branch.objects.filter(status='active')
    return render(request, 'operations/sales_list.html', {
        'sales': sales_qs,
        'branches': branches,
        'is_admin': is_admin
    })

@login_required
def expense_head_list_view(request):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    expense_heads = ExpenseHead.objects.all()
    return render(request, 'expense_heads/expense_head_list.html', {'expense_heads': expense_heads})


@login_required
def expense_head_create_view(request):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()

        if not name:
            messages.error(request, "Expense head name is required.")
        elif ExpenseHead.objects.filter(name__iexact=name).exists():
            messages.error(request, "This expense head already exists.")
        else:
            ExpenseHead.objects.create(name=name)
            messages.success(request, f"Expense head '{name}' created successfully.")
            return redirect('expense_head_list')

    return render(request, 'expense_heads/expense_head_form.html', {'title': 'Add New Expense Head'})


@login_required
def expense_head_delete_view(request, pk):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    expense_head = get_object_or_404(ExpenseHead, pk=pk)
    name = expense_head.name
    expense_head.delete()
    messages.success(request, f"Expense head '{name}' deleted successfully.")
    return redirect('expense_head_list')

@login_required
def expense_edit_view(request, pk):
    from operations.models import Expense

    expense = get_object_or_404(Expense, pk=pk)

    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    if not is_admin and user_profile and expense.branch != user_profile.branch:
        messages.error(request, "Permission denied.")
        return redirect('expense_list')

    branches = Branch.objects.filter(status='active')
    expense_heads = ExpenseHead.objects.all()

    if request.method == 'POST':
        branch_id = request.POST.get('branch')
        expense_date = request.POST.get('expense_date')
        expense_head_id = request.POST.get('expense_head')
        amount = request.POST.get('amount')
        description = request.POST.get('description', '').strip()

        if not all([branch_id, expense_date, expense_head_id, amount]):
            messages.error(request, "Please fill all required fields.")
        else:
            expense.branch_id = branch_id
            expense.expense_date = expense_date
            expense.expense_head_id = expense_head_id
            expense.amount = amount
            expense.description = description
            expense.updated_by = request.user
            expense.save()

            messages.success(request, "Expense updated successfully.")
            return redirect('expense_list')

    return render(request, 'operations/expense_form.html', {
        'expense': expense,
        'branches': branches,
        'expense_heads': expense_heads,
        'is_admin': is_admin,
        'title': 'Edit Expense',
    })


@login_required
def vehicle_list_view(request):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    vehicles = Vehicle.objects.select_related('branch').all()
    return render(request, 'vehicles/vehicle_list.html', {'vehicles': vehicles})


@login_required
def vehicle_create_view(request):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    branches = Branch.objects.all()

    if request.method == 'POST':
        vehicle_number = request.POST.get('vehicle_number', '').strip()
        vehicle_type = request.POST.get('vehicle_type', '').strip()
        branch_id = request.POST.get('branch')
        driver_name = request.POST.get('driver_name', '').strip()
        driver_phone = request.POST.get('driver_phone', '').strip()

        if not all([vehicle_number, vehicle_type, branch_id, driver_name, driver_phone]):
            messages.error(request, "All fields are required.")
        elif Vehicle.objects.filter(vehicle_number__iexact=vehicle_number).exists():
            messages.error(request, "This vehicle number already exists.")
        else:
            Vehicle.objects.create(
                vehicle_number=vehicle_number,
                vehicle_type=vehicle_type,
                branch_id=branch_id,
                driver_name=driver_name,
                driver_phone=driver_phone
            )
            messages.success(request, "Vehicle created successfully.")
            return redirect('vehicle_list')

    return render(request, 'vehicles/vehicle_form.html', {
        'title': 'Add New Vehicle',
        'branches': branches
    })

@login_required
def vehicle_edit_view(request, pk):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    vehicle = get_object_or_404(Vehicle, pk=pk)
    branches = Branch.objects.all()

    if request.method == 'POST':
        vehicle_number = request.POST.get('vehicle_number', '').strip()
        vehicle_type = request.POST.get('vehicle_type', '').strip()
        branch_id = request.POST.get('branch')
        driver_name = request.POST.get('driver_name', '').strip()
        driver_phone = request.POST.get('driver_phone', '').strip()

        if not all([vehicle_number, vehicle_type, branch_id, driver_name, driver_phone]):
            messages.error(request, "All fields are required.")
        elif Vehicle.objects.filter(
            vehicle_number__iexact=vehicle_number
        ).exclude(pk=pk).exists():
            messages.error(request, "This vehicle number already exists.")
        else:
            vehicle.vehicle_number = vehicle_number
            vehicle.vehicle_type = vehicle_type
            vehicle.branch_id = branch_id
            vehicle.driver_name = driver_name
            vehicle.driver_phone = driver_phone
            vehicle.save()

            messages.success(request, "Vehicle updated successfully.")
            return redirect('vehicle_list')

    return render(request, 'vehicles/vehicle_form.html', {
        'title': 'Edit Vehicle',
        'vehicle': vehicle,
        'branches': branches
    })


@login_required
def vehicle_delete_view(request, pk):
    if not request.user.profile.is_admin:
        messages.error(request, "Permission denied.")
        return redirect('dashboard')

    vehicle = get_object_or_404(Vehicle, pk=pk)
    vehicle.delete()

    messages.success(request, "Vehicle deleted successfully.")
    return redirect

@login_required
def expense_list_view(request):
    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    expenses = Expense.objects.select_related(
        'branch',
        'expense_head',
        'staff'
    ).all().order_by('-expense_date', '-id')

    if not is_admin and user_profile and user_profile.branch:
        expenses = expenses.filter(branch=user_profile.branch)
    else:
        branch_id = request.GET.get('branch_id')
        if branch_id:
            expenses = expenses.filter(branch_id=branch_id)

    expense_head_id = request.GET.get('expense_head_id')
    if expense_head_id:
        expenses = expenses.filter(expense_head_id=expense_head_id)

    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    if date_from:
        expenses = expenses.filter(expense_date__gte=date_from)

    if date_to:
        expenses = expenses.filter(expense_date__lte=date_to)

    total_expense = expenses.aggregate(
        total=Sum('amount')
    )['total'] or 0

    branches = Branch.objects.filter(status='active')
    expense_heads = ExpenseHead.objects.all()

    return render(request, 'operations/expense_list.html', {
        'expenses': expenses,
        'branches': branches,
        'expense_heads': expense_heads,
        'is_admin': is_admin,
        'date_from': date_from or '',
        'date_to': date_to or '',
        'expense_head_id': expense_head_id or '',
        'total_expense': total_expense,
    })


@login_required
def expense_export_excel_view(request):
    expenses = Expense.objects.select_related('branch', 'expense_head', 'staff').all()

    branch_id = request.GET.get('branch_id')
    if branch_id:
        expenses = expenses.filter(branch_id=branch_id)

    expense_head_id = request.GET.get('expense_head_id')
    if expense_head_id:
        expenses = expenses.filter(expense_head_id=expense_head_id)

    date_from = request.GET.get('date_from')
    if date_from:
        expenses = expenses.filter(expense_date__gte=date_from)

    date_to = request.GET.get('date_to')
    if date_to:
        expenses = expenses.filter(expense_date__lte=date_to)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Expenses'

    sheet.append(['Date', 'Branch', 'Expense Head', 'Amount', 'Description', 'Added By'])

    for expense in expenses:
        staff_name = expense.staff.username if expense.staff else ''
        sheet.append([
            expense.expense_date,
            expense.branch.name,
            expense.expense_head.name,
            float(expense.amount),
            expense.description or '',
            staff_name,
        ])

    for cell in sheet['A'][1:]:
        cell.number_format = 'DD-MM-YYYY'

    from openpyxl.styles import Alignment

    for row in sheet.iter_rows():
     for cell in row:
         cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

    # Center align and wrap text
    for row in sheet.iter_rows():
        for cell in row:
          cell.alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=True
        )


    sheet.column_dimensions['A'].width = 15
    sheet.column_dimensions['B'].width = 20
    sheet.column_dimensions['C'].width = 25
    sheet.column_dimensions['D'].width = 15
    sheet.column_dimensions['E'].width = 50
    sheet.column_dimensions['F'].width = 20

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename="expense_report.xlsx"'

    workbook.save(response)

    return response

@login_required
def expense_vehicle_details_view(request):
    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    branch_id = request.GET.get('branch_id')

    vehicles = Vehicle.objects.select_related('branch').all()

    if not is_admin and user_profile and user_profile.branch:
        vehicles = vehicles.filter(branch=user_profile.branch)

    elif branch_id:
        vehicles = vehicles.filter(branch_id=branch_id)

    branches = Branch.objects.filter(status='active')

    return render(request, 'operations/expense_vehicle_details.html', {
        'vehicles': vehicles,
        'branches': branches,
        'is_admin': is_admin,
        'selected_branch_id': branch_id or '',
    })

@login_required
def expense_entry_view(request):
    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    today = timezone.now().date()

    # ---------------------------------
    # Active branches
    # ---------------------------------

    if not is_admin and user_profile and user_profile.branch:
        branches = Branch.objects.filter(
            id=user_profile.branch.id,
            status='active'
        )
    else:
        branches = Branch.objects.filter(status='active')

    # ---------------------------------
    # Expense heads
    # ---------------------------------

    expense_heads = ExpenseHead.objects.all()

    # ---------------------------------
    # Vehicles
    # ---------------------------------

    vehicles = Vehicle.objects.select_related('branch').all()

    # ---------------------------------
    # GET values
    # ---------------------------------

    selected_branch_id = request.GET.get('branch_id')
    selected_expense_head_id = request.GET.get('expense_head_id')

    entry_mode = request.GET.get('entry_mode', '').strip()

    selected_branch = None
    selected_expense_head = None

    # ---------------------------------
    # Selected branch
    # ---------------------------------

    if selected_branch_id:

        if not is_admin and user_profile and user_profile.branch:

            if str(selected_branch_id) != str(user_profile.branch.id):
                messages.error(request, "Permission denied.")
                return redirect('expense_entry')

        selected_branch = Branch.objects.filter(
            id=selected_branch_id,
            status='active'
        ).first()

    # ---------------------------------
    # Selected expense head
    # ---------------------------------

    if selected_expense_head_id:

        selected_expense_head = ExpenseHead.objects.filter(
            id=selected_expense_head_id
        ).first()

    # ---------------------------------
    # IMPORTANT
    # Do NOT auto select Branch
    # ---------------------------------

    if entry_mode == 'branch' and selected_branch:

        entry_mode = 'branch'

    elif entry_mode == 'expense_head' and selected_expense_head:

        entry_mode = 'expense_head'

    else:

        entry_mode = ''

    # ---------------------------------
    # Branch mode rows
    # ---------------------------------

    branch_expense_rows = []

    if entry_mode == 'branch' and selected_branch:

        for head in expense_heads:

            if head.name.strip().lower() == 'vehicle expense':

                vehicle_list = vehicles.filter(
                    branch=selected_branch
                )

            else:

                vehicle_list = Vehicle.objects.none()

            branch_expense_rows.append({
                'expense_head': head,
                'vehicle_list': vehicle_list
            })

    # ---------------------------------
    # Expense head mode rows
    # ---------------------------------

    branch_rows = []

    if entry_mode == 'expense_head' and selected_expense_head:

        for branch in branches:

            branch_vehicles = vehicles.filter(
                branch=branch
            )

            branch_rows.append({
                'branch': branch,
                'vehicles': branch_vehicles
            })

    # ---------------------------------
    # POST
    # ---------------------------------

    if request.method == 'POST':

        post_mode = request.POST.get('entry_mode', '').strip()

        expense_date = request.POST.get('expense_date', '').strip()

        if not expense_date:

            messages.error(
                request,
                "Please select the expense date."
            )

            return redirect('expense_entry')

        saved_count = 0

        # =========================================
        # BRANCH MODE
        # =========================================

        if post_mode == 'branch':

            branch_id = request.POST.get('branch_id')

            if not branch_id:

                messages.error(
                    request,
                    "Please select a branch."
                )

                return redirect('expense_entry')

            # Permission check
            if (
                not is_admin
                and user_profile
                and user_profile.branch
                and str(branch_id) != str(user_profile.branch.id)
            ):

                messages.error(
                    request,
                    "Permission denied."
                )

                return redirect('expense_entry')

            branch = get_object_or_404(
                Branch,
                id=branch_id,
                status='active'
            )

            for head in expense_heads:

                amount = request.POST.get(
                    f'amount_{head.id}',
                    ''
                ).strip()

                remark = request.POST.get(
                    f'remark_{head.id}',
                    ''
                ).strip()

                vehicle_id = request.POST.get(
                    f'vehicle_{head.id}'
                )

                # Nothing entered
                if not amount and not remark:

                    continue

                # Amount required
                if not amount:

                    messages.error(
                        request,
                        f"Amount is required for {head.name}."
                    )

                    return redirect(
                        f'/expenses/add/?branch_id={branch.id}&entry_mode=branch'
                    )

                description = remark

                # ---------------------------------
                # Vehicle expense
                # ---------------------------------

                if (
                    head.name.strip().lower()
                    == 'vehicle expense'
                    and vehicle_id
                ):

                    try:

                        vehicle = Vehicle.objects.select_related(
                            'branch'
                        ).get(
                            id=vehicle_id,
                            branch=branch
                        )

                        vehicle_details = (
                            f"Vehicle Number: "
                            f"{vehicle.vehicle_number}\n"
                            f"Vehicle Type: "
                            f"{vehicle.vehicle_type}\n"
                            f"Branch: "
                            f"{vehicle.branch.name}\n"
                            f"Driver Name: "
                            f"{vehicle.driver_name}\n"
                            f"Driver Phone: "
                            f"{vehicle.driver_phone}"
                        )

                        if description:

                            description = (
                                f"{description}\n\n"
                                f"{vehicle_details}"
                            )

                        else:

                            description = vehicle_details

                    except Vehicle.DoesNotExist:

                        messages.error(
                            request,
                            "Selected vehicle not found for the selected branch."
                        )

                        return redirect(
                            f'/expenses/add/?branch_id={branch.id}&entry_mode=branch'
                        )

                # ---------------------------------
                # Save expense
                # ---------------------------------

                Expense.objects.create(
                    branch=branch,
                    staff=request.user,
                    expense_date=expense_date,
                    expense_head=head,
                    amount=amount,
                    description=description,
                    created_by=request.user,
                    updated_by=request.user
                )

                saved_count += 1

        # =========================================
        # EXPENSE HEAD MODE
        # =========================================

        elif post_mode == 'expense_head':

            expense_head_id = request.POST.get(
                'expense_head_id'
            )

            if not expense_head_id:

                messages.error(
                    request,
                    "Please select an expense head."
                )

                return redirect('expense_entry')

            expense_head = get_object_or_404(
                ExpenseHead,
                id=expense_head_id
            )

            for branch in branches:

                # Permission check
                if (
                    not is_admin
                    and user_profile
                    and user_profile.branch
                    and branch.id != user_profile.branch.id
                ):

                    continue

                amount = request.POST.get(
                    f'branch_amount_{branch.id}',
                    ''
                ).strip()

                remark = request.POST.get(
                    f'branch_remark_{branch.id}',
                    ''
                ).strip()

                vehicle_id = request.POST.get(
                    f'branch_vehicle_{branch.id}'
                )

                # Nothing entered
                if not amount and not remark:

                    continue

                # Amount required
                if not amount:

                    messages.error(
                        request,
                        f"Amount is required for {branch.name}."
                    )

                    return redirect(
                        f'/expenses/add/?expense_head_id={expense_head.id}&entry_mode=expense_head'
                    )

                description = remark

                # ---------------------------------
                # Vehicle expense
                # ---------------------------------

                if (
                    expense_head.name.strip().lower()
                    == 'vehicle expense'
                    and vehicle_id
                ):

                    try:

                        vehicle = Vehicle.objects.select_related(
                            'branch'
                        ).get(
                            id=vehicle_id,
                            branch=branch
                        )

                        vehicle_details = (
                            f"Vehicle Number: "
                            f"{vehicle.vehicle_number}\n"
                            f"Vehicle Type: "
                            f"{vehicle.vehicle_type}\n"
                            f"Branch: "
                            f"{vehicle.branch.name}\n"
                            f"Driver Name: "
                            f"{vehicle.driver_name}\n"
                            f"Driver Phone: "
                            f"{vehicle.driver_phone}"
                        )

                        if description:

                            description = (
                                f"{description}\n\n"
                                f"{vehicle_details}"
                            )

                        else:

                            description = vehicle_details

                    except Vehicle.DoesNotExist:

                        messages.error(
                            request,
                            "Selected vehicle not found for the selected branch."
                        )

                        return redirect(
                            f'/expenses/add/?expense_head_id={expense_head.id}&entry_mode=expense_head'
                        )

                # ---------------------------------
                # Save expense
                # ---------------------------------

                Expense.objects.create(
                    branch=branch,
                    staff=request.user,
                    expense_date=expense_date,
                    expense_head=expense_head,
                    amount=amount,
                    description=description,
                    created_by=request.user,
                    updated_by=request.user
                )

                saved_count += 1

        # =========================================
        # INVALID MODE
        # =========================================

        else:

            messages.error(
                request,
                "Please select Branch or Expense Head."
            )

            return redirect('expense_entry')

        # =========================================
        # SUCCESS
        # =========================================

        if saved_count > 0:

            messages.success(
                request,
                f"{saved_count} expense record(s) added successfully."
            )

            return redirect('expense_list')

        messages.warning(
            request,
            "No expense data entered."
        )

        return redirect('expense_entry')

    # ---------------------------------
    # Render
    # ---------------------------------

    return render(
        request,
        'operations/expense_form.html',
        {
            'branches': branches,
            'expense_heads': expense_heads,
            'vehicles': vehicles,
            'is_admin': is_admin,
            'selected_branch': selected_branch,
            'selected_expense_head': selected_expense_head,
            'entry_mode': entry_mode,
            'branch_expense_rows': branch_expense_rows,
            'branch_rows': branch_rows,
            'today': today,
            'title': 'Add Expense',
        }
    )

@login_required
def expense_edit_view(request, pk):
    expense = get_object_or_404(
        Expense.objects.select_related(
            'branch',
            'expense_head'
        ),
        pk=pk
    )

    user_profile = getattr(request.user, 'profile', None)
    is_admin = (
        user_profile.is_admin
        if user_profile
        else request.user.is_superuser
    )

    # Non-admin can edit only own branch expense
    if (
        not is_admin
        and user_profile
        and user_profile.branch
        and expense.branch_id != user_profile.branch.id
    ):
        messages.error(request, "Permission denied.")
        return redirect('expense_list')

    if is_admin:
        branches = Branch.objects.filter(
            status='active'
        )
    else:
        branches = Branch.objects.filter(
            id=user_profile.branch.id,
            status='active'
        )

    expense_heads = ExpenseHead.objects.all()

    if request.method == 'POST':
        branch_id = request.POST.get(
            'branch'
        )

        expense_date = request.POST.get(
            'expense_date'
        )

        expense_head_id = request.POST.get(
            'expense_head'
        )

        amount = request.POST.get(
            'amount'
        )

        description = request.POST.get(
            'description',
            ''
        ).strip()

        if not all([
            branch_id,
            expense_date,
            expense_head_id,
            amount
        ]):
            messages.error(
                request,
                "Please fill all required fields."
            )

        else:
            # Non-admin cannot move expense to another branch
            if (
                not is_admin
                and user_profile
                and user_profile.branch
                and str(branch_id)
                != str(user_profile.branch.id)
            ):
                messages.error(
                    request,
                    "Permission denied."
                )
                return redirect('expense_list')

            branch = get_object_or_404(
                Branch,
                id=branch_id,
                status='active'
            )

            expense_head = get_object_or_404(
                ExpenseHead,
                id=expense_head_id
            )

            expense.branch = branch
            expense.expense_date = expense_date
            expense.expense_head = expense_head
            expense.amount = amount
            expense.description = description
            expense.updated_by = request.user

            expense.save()

            messages.success(
                request,
                "Expense updated successfully."
            )

            return redirect('expense_list')

    return render(
        request,
        'operations/expense_edit_form.html',
        {
            'expense': expense,
            'branches': branches,
            'expense_heads': expense_heads,
            'is_admin': is_admin,
            'title': 'Edit Expense',
        }
    )
@login_required
def expense_delete_view(request, pk):
    from operations.models import Expense

    expense = get_object_or_404(Expense, pk=pk)

    user_profile = getattr(request.user, 'profile', None)
    is_admin = user_profile.is_admin if user_profile else request.user.is_superuser

    if not is_admin and user_profile and expense.branch != user_profile.branch:
        messages.error(request, "Permission denied.")
        return redirect('expense_list')

    expense.delete()
    messages.success(request, "Expense deleted successfully.")
    return redirect('expense_list')